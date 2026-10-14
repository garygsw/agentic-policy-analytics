"""Generate the mock internal workbook: synthetic tech-talent programme records of a fictional agency.

    uv run python scripts/make_internal_mock.py

The numbers are invented (fixed random seed, so the file is reproducible). The data contains a few problems
on purpose, so the cleaning step has something to catch: a "Total" row per year, a duplicate row, a missing
value written "n.a." and a text value "TBC".
"""
import random
from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parents[1] / "data" / "internal" / "tech_talent_programmes.xlsx"
PROGRAMMES = {"Mid-career conversion": 1.0, "Graduate traineeship": 0.8, "Industry upskilling": 1.6}
TRACKS = {"Software engineering": 1.0, "Data and AI": 0.7, "Cybersecurity": 0.45, "Cloud and infrastructure": 0.5}
YEARS = range(2018, 2026)


def rows(rng: random.Random) -> list[list]:
    out = []
    for year in YEARS:
        growth = 1 + 0.09 * (year - 2018) - (0.12 if year == 2020 else 0)  # dip in 2020, steady growth after
        year_rows = []
        for programme, p_scale in PROGRAMMES.items():
            for track, t_scale in TRACKS.items():
                participants = round(400 * p_scale * t_scale * growth * rng.uniform(0.9, 1.1))
                placed = round(participants * rng.uniform(0.62, 0.86))
                salary = round(4200 * (1 + 0.035 * (year - 2018)) * (1.15 if track == "Data and AI" else 1.0), -1)
                year_rows.append([year, programme, track, participants, placed, salary])
        out += year_rows
        out.append([year, "Total", "Total", sum(r[3] for r in year_rows), sum(r[4] for r in year_rows), None])
    out.append(list(out[5]))  # duplicate record
    out[10][4] = "n.a."       # missing value
    out[20][5] = "TBC"        # not a number
    return out


def main() -> None:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["SYNTHETIC DATA - for demonstrating the platform only. Not an official statistic."])
    readme.append(["Fictional agency records of tech-talent programme intakes and placements, 2018-2025."])
    sheet = workbook.create_sheet("placements")
    sheet.append(["year", "programme", "track", "participants", "placed_within_6_months",
                  "median_starting_salary_sgd"])
    for row in rows(random.Random(2026)):
        sheet.append(row)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)
    print(f"Wrote {OUT.relative_to(OUT.parents[2])}: {sheet.max_row - 1} rows")


if __name__ == "__main__":
    main()
